import pytest
from openpyxl import Workbook

from app.services.profiling import DatasetFileError, inspect_upload, profile_file


def test_csv_profile_handles_bom_delimiter_duplicate_headers_missing_and_duplicates(tmp_path):
    path = tmp_path / "source.csv"
    path.write_text(
        "\ufeffValue;Value;Note;Flag\n1;10;alpha;true\n;20;beta;false\n1;10;alpha;true\n",
        encoding="utf-8",
    )
    profile = profile_file(path, "csv", None, max_rows=100, max_columns=10)

    assert (profile["row_count"], profile["column_count"]) == (3, 4)
    assert profile["duplicate_row_count"] == 1
    assert profile["quality_summary"]["missing_cell_count"] == 1
    assert [column["id"] for column in profile["columns"][:2]] == ["column_001", "column_002"]
    assert [column["name"] for column in profile["columns"][:2]] == ["Value", "Value"]
    assert profile["columns"][0]["data_type"] == "integer"
    assert profile["columns"][0]["numeric_summary"] == {
        "min": 1.0,
        "max": 1.0,
        "mean": 1.0,
        "median": 1.0,
        "median_is_approximate": False,
        "non_finite_count": 0,
    }
    assert profile["columns"][3]["data_type"] == "boolean"


def test_xlsx_requires_a_sheet_choice_and_profiles_only_the_selected_sheet(tmp_path):
    path = tmp_path / "workbook.xlsx"
    book = Workbook()
    first = book.active
    first.title = "Sales"
    first.append(["Amount", "Region"])
    first.append([2.0, "West"])
    first.append([None, "West"])
    second = book.create_sheet("Notes")
    second.append(["Text"])
    second.append(["Do not merge"])
    book.save(path)

    with pytest.raises(DatasetFileError, match="Choose one worksheet") as error:
        inspect_upload(path, "xlsx", None, 10)
    assert error.value.details == {"sheets": ["Sales", "Notes"]}

    profile = profile_file(path, "xlsx", "Sales", max_rows=100, max_columns=10)
    assert profile["selected_sheet"] == "Sales"
    assert (profile["row_count"], profile["column_count"]) == (2, 2)
    assert profile["columns"][0]["numeric_summary"]["mean"] == 2.0
    assert profile["columns"][1]["distinct_count"] == 1


def test_invalid_utf8_formulas_and_row_bounds_are_rejected(tmp_path):
    invalid = tmp_path / "invalid.csv"
    invalid.write_bytes(b"name\n\xff")
    with pytest.raises(DatasetFileError, match="UTF-8"):
        inspect_upload(invalid, "csv", None, 10)

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["value"])
    sheet.append(["=1+1"])
    formula = tmp_path / "formula.xlsx"
    workbook.save(formula)
    with pytest.raises(DatasetFileError, match="Formula cells"):
        profile_file(formula, "xlsx", "Sheet", max_rows=100, max_columns=10)

    rows = tmp_path / "rows.csv"
    rows.write_text("value\n1\n2\n3\n", encoding="utf-8")
    with pytest.raises(DatasetFileError, match="row limit"):
        profile_file(rows, "csv", None, max_rows=2, max_columns=10)
