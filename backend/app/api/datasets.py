"""Authenticated dataset upload, status, profile, version, and preview APIs."""

import hashlib
import os
import re
import tempfile
from datetime import date, datetime
from math import isfinite
from pathlib import Path
from typing import Annotated
from uuid import UUID, uuid4

import pandas as pd
from fastapi import APIRouter, Depends, File, Form, Query, Request, UploadFile
from openpyxl import load_workbook
from pydantic import BaseModel
from rq import Retry
from sqlalchemy import func, select
from sqlalchemy.orm import Session
from starlette.concurrency import run_in_threadpool

from app.api.auth import DB, CurrentUser, Envelope, result
from app.core.errors import AppError
from app.core.security import require_csrf, require_upload_csrf
from app.db.models import Dataset, DatasetProfile, DatasetVersion, StoredArtifact
from app.services.profiling import DatasetFileError, inspect_upload
from app.storage.paths import ArtifactStore, cleanup_abandoned_uploads
from worker.connection import queue

router = APIRouter(prefix="/api/v1/datasets", tags=["datasets"])
CHUNK_SIZE = 1024 * 1024
FILENAME = re.compile(r"[^\x00-\x1f\x7f/\\]{1,255}\Z")


class DatasetListItem(BaseModel):
    id: UUID
    display_name: str
    original_filename: str
    source_format: str
    file_size_bytes: int
    created_at: datetime
    status: str
    selected_sheet: str | None = None
    row_count: int | None = None
    column_count: int | None = None


class DatasetCollection(BaseModel):
    items: list[DatasetListItem]
    page: int
    page_size: int
    total: int


class NumericSummary(BaseModel):
    min: float
    max: float
    mean: float
    median: float
    median_is_approximate: bool
    non_finite_count: int


class ProfileColumn(BaseModel):
    id: str
    name: str
    data_type: str
    missing_count: int
    missing_percent: float
    distinct_count: int
    distinct_count_is_exact: bool
    sample_values: list[str | int | float | bool]
    numeric_summary: NumericSummary | None


class ProfileQuality(BaseModel):
    missing_cell_count: int
    missing_cell_percent: float
    columns_with_missing_values: int
    duplicate_row_count: int


class DatasetProfileView(BaseModel):
    row_count: int
    column_count: int
    columns: list[ProfileColumn]
    duplicate_row_count: int
    quality_summary: ProfileQuality
    methodology: dict[str, str]
    selected_sheet: str | None


class DatasetVersionView(BaseModel):
    id: UUID
    version_no: int
    kind: str
    created_at: datetime
    file_size_bytes: int
    checksum: str
    mime_type: str


class VersionReference(BaseModel):
    id: UUID
    number: int
    kind: str


class DatasetStatusView(DatasetListItem):
    processing_error: str | None = None


class DatasetDetailView(DatasetListItem):
    checksum: str | None = None
    version: VersionReference | None = None
    processing_error: str | None = None


class DatasetPreviewView(BaseModel):
    columns: list[ProfileColumn]
    rows: list[dict[str, str | int | float | bool | None]]
    page: int
    page_size: int
    total_rows: int
    total_pages: int


def _owned_dataset(db: Session, owner_id: UUID, dataset_id: UUID) -> Dataset:
    dataset = db.scalar(
        select(Dataset).where(Dataset.id == dataset_id, Dataset.owner_id == owner_id)
    )
    if dataset is None:
        raise AppError(404, "DATASET_NOT_FOUND", "Dataset not found.")
    return dataset


def _profile(db: Session, dataset: Dataset):
    version = db.scalar(
        select(DatasetVersion).where(
            DatasetVersion.dataset_id == dataset.id,
            DatasetVersion.owner_id == dataset.owner_id,
            DatasetVersion.version_no == 1,
        )
    )
    if version is None:
        return None, None
    return version, db.scalar(
        select(DatasetProfile).where(DatasetProfile.dataset_version_id == version.id)
    )


def _ready(dataset: Dataset):
    if dataset.status == "failed":
        raise AppError(
            409,
            "DATASET_PROCESSING_FAILED",
            dataset.processing_error or "Dataset processing failed.",
        )
    if dataset.status != "ready":
        raise AppError(409, "DATASET_NOT_READY", "Dataset profiling is still in progress.")


@router.post(
    "",
    status_code=201,
    dependencies=[Depends(require_upload_csrf)],
    response_model=Envelope[DatasetListItem],
)
async def upload_dataset(
    request: Request,
    user: CurrentUser,
    db: DB,
    file: Annotated[UploadFile, File()],
    sheet_name: Annotated[str | None, Form(max_length=100)] = None,
):
    settings = request.app.state.settings
    filename = file.filename or ""
    if not FILENAME.fullmatch(filename) or filename in {".", ".."}:
        raise AppError(422, "INVALID_FILENAME", "Choose a file with a simple filename.")
    suffix = Path(filename).suffix.casefold()
    if suffix not in {".csv", ".xlsx"}:
        raise AppError(415, "UNSUPPORTED_FILE", "Only CSV and XLSX files are supported.")
    source_format = suffix[1:]
    media_type = (
        "text/csv"
        if source_format == "csv"
        else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    root = Path(settings.file_storage_root)
    root.mkdir(parents=True, exist_ok=True)
    store = ArtifactStore(root)
    cleanup_abandoned_uploads(root)
    temporary = None
    finalized = False
    reference = f"{uuid4().hex}.{source_format}"
    stored = store.resolve(reference)
    size = 0
    checksum = hashlib.sha256()
    try:
        with tempfile.NamedTemporaryFile(
            prefix=".upload-", suffix=".part", dir=root, delete=False
        ) as stream:
            temporary = Path(stream.name)
            while chunk := await file.read(CHUNK_SIZE):
                size += len(chunk)
                if size > settings.max_upload_bytes:
                    limit_mb = settings.max_upload_bytes / 1_000_000
                    raise AppError(
                        413,
                        "UPLOAD_TOO_LARGE",
                        f"The file exceeds the {limit_mb:g} MB upload limit.",
                    )
                checksum.update(chunk)
                stream.write(chunk)
        if size == 0:
            raise AppError(422, "INVALID_FILE", "The uploaded file is empty.")
        try:
            file_info = await run_in_threadpool(
                inspect_upload, temporary, source_format, sheet_name, settings.max_dataset_columns
            )
        except DatasetFileError as exc:
            raise AppError(422, exc.code, str(exc), exc.details) from None
        os.link(temporary, stored)
        finalized = True
        temporary.unlink()
        temporary = None
        dataset = Dataset(
            owner_id=user.id,
            display_name=Path(filename).stem[:200] or "Uploaded dataset",
            source_format=source_format,
            original_filename=filename,
            media_type=media_type,
            file_size_bytes=size,
            status="queued",
            selected_sheet=file_info["sheet_name"],
        )
        artifact = StoredArtifact(
            owner_id=user.id,
            relative_path=reference,
            mime_type=media_type,
            size_bytes=size,
            sha256=checksum.hexdigest(),
        )
        db.add_all([dataset, artifact])
        db.flush()
        version = DatasetVersion(
            owner_id=user.id,
            dataset_id=dataset.id,
            artifact_id=artifact.id,
            version_no=1,
            kind="raw",
        )
        db.add(version)
        db.commit()
    except AppError:
        db.rollback()
        if finalized:
            stored.unlink(missing_ok=True)
        raise
    except Exception:
        db.rollback()
        if finalized:
            stored.unlink(missing_ok=True)
        raise AppError(
            503, "UPLOAD_FAILED", "The upload could not be saved. Please try again."
        ) from None
    finally:
        await file.close()
        if temporary:
            try:
                temporary.unlink(missing_ok=True)
            except OSError:
                pass
    try:
        redis = request.app.state.redis
        if redis is None:
            raise RuntimeError("queue unavailable")
        queue(settings, redis).enqueue(
            "worker.jobs.profile_dataset",
            str(dataset.id),
            job_id=f"profile-{dataset.id}",
            job_timeout=min(settings.max_job_seconds, 180),
            result_ttl=3600,
            failure_ttl=86400,
            retry=Retry(max=1),
        )
    except Exception:
        db.refresh(dataset)
        dataset.status = "failed"
        dataset.processing_error = "The profiling queue is unavailable. Retry profiling."
        db.commit()
        raise AppError(
            503,
            "QUEUE_UNAVAILABLE",
            "The upload was saved but profiling could not start. Retry it from the dataset.",
        ) from None
    return result(request, _list_item(db, dataset))


def _list_item(db: Session, dataset: Dataset):
    version, profile = _profile(db, dataset)
    payload = {
        "id": dataset.id,
        "display_name": dataset.display_name,
        "original_filename": dataset.original_filename,
        "source_format": dataset.source_format,
        "file_size_bytes": dataset.file_size_bytes,
        "created_at": dataset.created_at,
        "status": dataset.status,
        "selected_sheet": dataset.selected_sheet,
        "row_count": profile.row_count if profile else None,
        "column_count": profile.column_count if profile else None,
    }
    return DatasetListItem.model_validate(payload).model_dump(mode="json")


@router.get("", response_model=Envelope[DatasetCollection])
def list_datasets(
    request: Request,
    user: CurrentUser,
    db: DB,
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
):
    total = (
        db.scalar(select(func.count()).select_from(Dataset).where(Dataset.owner_id == user.id)) or 0
    )
    datasets = db.scalars(
        select(Dataset)
        .where(Dataset.owner_id == user.id)
        .order_by(Dataset.created_at.desc(), Dataset.id)
        .offset((page - 1) * page_size)
        .limit(page_size)
    ).all()
    return result(
        request,
        {
            "items": [_list_item(db, item) for item in datasets],
            "page": page,
            "page_size": page_size,
            "total": total,
        },
    )


@router.get("/{dataset_id}/status", response_model=Envelope[DatasetStatusView])
def dataset_status(request: Request, dataset_id: UUID, user: CurrentUser, db: DB):
    dataset = _owned_dataset(db, user.id, dataset_id)
    payload = _list_item(db, dataset)
    payload["processing_error"] = dataset.processing_error
    return result(request, payload)


@router.post(
    "/{dataset_id}/retry",
    dependencies=[Depends(require_csrf)],
    response_model=Envelope[DatasetListItem],
)
def retry_profile(request: Request, dataset_id: UUID, user: CurrentUser, db: DB):
    dataset = db.scalar(
        select(Dataset)
        .where(Dataset.id == dataset_id, Dataset.owner_id == user.id)
        .with_for_update()
    )
    if dataset is None:
        raise AppError(404, "DATASET_NOT_FOUND", "Dataset not found.")
    if dataset.status != "failed":
        raise AppError(409, "DATASET_NOT_FAILED", "Only a failed dataset can be retried.")
    dataset.status = "queued"
    dataset.processing_error = None
    db.commit()
    try:
        if request.app.state.redis is None:
            raise RuntimeError("queue unavailable")
        queue(request.app.state.settings, request.app.state.redis).enqueue(
            "worker.jobs.profile_dataset",
            str(dataset.id),
            job_id=f"profile-{dataset.id}-{uuid4().hex}",
            job_timeout=min(request.app.state.settings.max_job_seconds, 180),
            result_ttl=3600,
            failure_ttl=86400,
            retry=Retry(max=1),
        )
    except Exception:
        dataset.status = "failed"
        dataset.processing_error = "The profiling queue is unavailable. Retry profiling."
        db.commit()
        raise AppError(503, "QUEUE_UNAVAILABLE", dataset.processing_error) from None
    return result(request, _list_item(db, dataset))


@router.get("/{dataset_id}", response_model=Envelope[DatasetDetailView])
def get_dataset(request: Request, dataset_id: UUID, user: CurrentUser, db: DB):
    dataset = _owned_dataset(db, user.id, dataset_id)
    version, profile = _profile(db, dataset)
    payload = _list_item(db, dataset)
    payload["checksum"] = None
    if version is not None:
        artifact = db.scalar(
            select(StoredArtifact).where(
                StoredArtifact.id == version.artifact_id,
                StoredArtifact.owner_id == user.id,
            )
        )
        payload["checksum"] = artifact.sha256 if artifact else None
        payload["version"] = {"id": version.id, "number": version.version_no, "kind": version.kind}
    payload["processing_error"] = dataset.processing_error
    return result(request, payload)


@router.get("/{dataset_id}/versions", response_model=Envelope[list[DatasetVersionView]])
def list_versions(request: Request, dataset_id: UUID, user: CurrentUser, db: DB):
    dataset = _owned_dataset(db, user.id, dataset_id)
    versions = db.execute(
        select(DatasetVersion, StoredArtifact)
        .join(StoredArtifact, StoredArtifact.id == DatasetVersion.artifact_id)
        .where(
            DatasetVersion.dataset_id == dataset.id,
            DatasetVersion.owner_id == user.id,
            StoredArtifact.owner_id == user.id,
        )
        .order_by(DatasetVersion.version_no)
    ).all()
    return result(
        request,
        [
            {
                "id": version.id,
                "version_no": version.version_no,
                "kind": version.kind,
                "created_at": version.created_at,
                "file_size_bytes": artifact.size_bytes,
                "checksum": artifact.sha256,
                "mime_type": artifact.mime_type,
            }
            for version, artifact in versions
        ],
    )


@router.get("/{dataset_id}/profile", response_model=Envelope[DatasetProfileView])
def get_profile(
    request: Request,
    dataset_id: UUID,
    user: CurrentUser,
    db: DB,
    version_id: UUID | None = None,
):
    dataset = _owned_dataset(db, user.id, dataset_id)
    _ready(dataset)
    version = db.scalar(
        select(DatasetVersion)
        .where(
            DatasetVersion.dataset_id == dataset.id,
            DatasetVersion.owner_id == user.id,
            *([DatasetVersion.id == version_id] if version_id else []),
        )
        .order_by(DatasetVersion.version_no.desc())
        .limit(1)
    )
    profile = (
        db.scalar(select(DatasetProfile).where(DatasetProfile.dataset_version_id == version.id))
        if version
        else None
    )
    if profile is None:
        raise AppError(409, "PROFILE_UNAVAILABLE", "Dataset profile is not available yet.")
    return result(request, profile.profile_json)


@router.get("/{dataset_id}/preview", response_model=Envelope[DatasetPreviewView])
def preview_dataset(
    request: Request,
    dataset_id: UUID,
    user: CurrentUser,
    db: DB,
    page: int = Query(1, ge=1),
    page_size: int = Query(50, ge=1, le=100),
    version_id: UUID | None = None,
):
    dataset = _owned_dataset(db, user.id, dataset_id)
    _ready(dataset)
    version = db.scalar(
        select(DatasetVersion)
        .where(
            DatasetVersion.dataset_id == dataset.id,
            DatasetVersion.owner_id == user.id,
            *([DatasetVersion.id == version_id] if version_id else []),
        )
        .order_by(DatasetVersion.version_no.desc())
        .limit(1)
    )
    profile = (
        db.scalar(select(DatasetProfile).where(DatasetProfile.dataset_version_id == version.id))
        if version
        else None
    )
    artifact = (
        db.scalar(
            select(StoredArtifact).where(
                StoredArtifact.id == version.artifact_id,
                StoredArtifact.owner_id == user.id,
            )
        )
        if version
        else None
    )
    if artifact is None or profile is None:
        raise AppError(409, "PREVIEW_UNAVAILABLE", "Dataset preview is not available yet.")
    path = ArtifactStore(Path(request.app.state.settings.file_storage_root)).resolve(
        artifact.relative_path
    )
    source_format = "xlsx" if artifact.mime_type.endswith("spreadsheetml.sheet") else "csv"
    sheet_name = dataset.selected_sheet if version.version_no == 1 else None
    info = inspect_upload(
        path,
        source_format,
        sheet_name,
        request.app.state.settings.max_dataset_columns,
    )
    offset = (page - 1) * page_size
    rows = _read_preview(path, source_format, sheet_name, info, offset, page_size)
    return result(
        request,
        {
            "columns": profile.profile_json["columns"],
            "rows": rows,
            "page": page,
            "page_size": page_size,
            "total_rows": profile.row_count,
            "total_pages": (profile.row_count + page_size - 1) // page_size,
        },
    )


def _json_value(value):
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if hasattr(value, "item"):
        value = value.item()
    if isinstance(value, float) and not isfinite(value):
        return None
    if isinstance(value, str) and value.strip().casefold() in {
        "",
        "na",
        "n/a",
        "nan",
        "null",
        "none",
    }:
        return None
    return value


def _read_preview(
    path: Path, source_format: str, sheet_name: str | None, info: dict, offset: int, limit: int
):
    names = [f"column_{index + 1:03d}" for index in range(len(info["headers"]))]
    if source_format == "csv":
        try:
            chunks = pd.read_csv(
                path,
                encoding="utf-8-sig",
                sep=info["delimiter"],
                names=names,
                header=0,
                dtype=object,
                on_bad_lines="error",
                skip_blank_lines=False,
                keep_default_na=False,
                na_values=["", "NA", "N/A", "NaN", "null", "none"],
                chunksize=2000,
            )
            remaining = offset
            for frame in chunks:
                if remaining >= len(frame):
                    remaining -= len(frame)
                    continue
                selected = frame.iloc[remaining : remaining + limit]
                return [
                    {name: _json_value(value) for name, value in zip(names, row, strict=True)}
                    for row in selected.itertuples(index=False, name=None)
                ]
            return []
        except Exception:
            raise AppError(
                500, "PREVIEW_FAILED", "The dataset preview could not be read."
            ) from None
    try:
        with path.open("rb") as stream:
            workbook = load_workbook(stream, read_only=True, data_only=False, keep_links=False)
            try:
                sheet = workbook[sheet_name]
                rows = []
                for row in sheet.iter_rows(
                    min_row=offset + 2, max_row=offset + limit + 1, values_only=False
                ):
                    values = []
                    for cell in row[: len(names)]:
                        if cell.data_type == "f":
                            raise AppError(
                                422, "FORMULAS_UNSUPPORTED", "Formula cells are not evaluated."
                            )
                        values.append(_json_value(cell.value))
                    values.extend([None] * (len(names) - len(values)))
                    rows.append(dict(zip(names, values, strict=True)))
            finally:
                workbook.close()
        return rows
    except AppError:
        raise
    except Exception:
        raise AppError(500, "PREVIEW_FAILED", "The dataset preview could not be read.") from None
